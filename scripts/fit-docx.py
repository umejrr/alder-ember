#!/usr/bin/env python3
"""Post-processes the pandoc Word copy of the brief: A4 page, readable tables that fill the page width with light borders.
    python3 scripts/fit-docx.py file.docx
"""
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

path = sys.argv[1]
doc = Document(path)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.0)
PAGE_CM = 17.0

def set_el(parent, tag, **attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el

for t in doc.tables:
    n = len(t.columns)
    weights = []
    for ci in range(n):
        texts = [r.cells[ci].text for r in t.rows]
        longest = max(len(x) for x in texts)
        word = max((len(w) for x in texts for w in x.split()), default=1)
        weights.append(max(min(longest, 55), word * 1.4 + 2, 4))  # words are weighted up: monospace and bold run wider than plain text
    total = sum(weights)
    cm = [PAGE_CM * w / total for w in weights]
    # never narrower than a digit plus the cell padding, or the text is clipped; the other columns give up the difference
    MIN = 1.1
    small = [i for i, w in enumerate(cm) if w < MIN]
    if small:
        spare = PAGE_CM - MIN * len(small)
        rest = sum(cm[i] for i in range(n) if i not in small)
        cm = [MIN if i in small else cm[i] * spare / rest for i in range(n)]
    widths = [Cm(w) for w in cm]
    pr = t._tbl.tblPr
    set_el(pr, 'w:tblW', **{'w:type': 'dxa', 'w:w': int(PAGE_CM / 2.54 * 1440)})
    set_el(pr, 'w:tblLayout', **{'w:type': 'fixed'})
    borders = set_el(pr, 'w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        set_el(borders, f'w:{side}', **{'w:val': 'single', 'w:sz': 4, 'w:space': 0, 'w:color': 'D8D2C8'})
    for ci, col in enumerate(t.columns):
        col.width = widths[ci]
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = widths[ci]
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
doc.save(path)
